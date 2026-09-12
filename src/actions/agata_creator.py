# src/actions/agata_creator.py
"""
Agata Creator - Generador de Documentos Word y Presentaciones PowerPoint
Soporte dual:
- Nube: Gemini (modelo gemini-3.5-flash-lite)
- Local: Ollama (modelo configurado en api_keys.json, ej. qwen2.5:3b)
Diseños profesionales con paletas de colores personalizables.
"""

import json
import os
import re
import subprocess
import sys
import time
import uuid
from pathlib import Path

import requests

# =====================================================================
# CONFIGURACIÓN
# =====================================================================

# Carpeta donde se guardarán los documentos generados
AGATA_FOLDER = Path.home() / "Documents" / "Agata_Projects"


def _ensure_folder():
    """Asegura que la carpeta de destino exista."""
    AGATA_FOLDER.mkdir(parents=True, exist_ok=True)


# =====================================================================
# PROVEEDORES DE IA (Nube y Local) - MEJORADOS
# =====================================================================

def _call_gemini(prompt: str, api_key: str = "", max_retries: int = 2) -> str:
    """
    Genera contenido usando Gemini (Nube).
    Modelo fijo: gemini-3.5-flash-lite (probado y funcional).
    Añade reintentos automáticos.
    """
    last_error = ""
    for attempt in range(max_retries):
        try:
            if not api_key:
                from src.core.config import _get_api_key
                api_key = _get_api_key()
            if not api_key:
                return "[Gemini Error]: No API key disponible"

            from google import genai
            client = genai.Client(api_key=api_key)
            response = client.models.generate_content(
                model="gemini-3.5-flash-lite",
                contents=prompt,
                config={"max_output_tokens": 2000}
            )
            if response and response.text:
                text = response.text.strip()
                if len(text) > 10:
                    return text
                last_error = "Respuesta vacía o demasiado corta"
            else:
                last_error = "Respuesta vacía"
        except Exception as e:
            last_error = str(e)[:200]
            if attempt < max_retries - 1:
                time.sleep(2)  # Esperar antes de reintentar
                continue
    return f"[Gemini Error]: {last_error}"


def _call_ollama(prompt: str, max_tokens: int = 4000) -> str:
    """
    Genera contenido usando Ollama (Local).
    Usa el modelo y la URL configurados en el sistema.
    Verifica conexión antes de intentar.
    """
    try:
        from src.core.config import get_config
        config = get_config()
        model = config.get("ollama_model", "qwen2.5:3b")
        base_url = config.get("ollama_url", "http://localhost:11434").rstrip("/")
        
        # Verificar si Ollama está corriendo
        try:
            health_check = requests.get(f"{base_url}/api/tags", timeout=3)
            if health_check.status_code != 200:
                return "[Ollama Error]: Ollama no responde correctamente"
        except requests.exceptions.ConnectionError:
            return "[Ollama Error]: No se pudo conectar a Ollama. ¿Está ejecutándose?"
        
        url = f"{base_url}/api/chat"
        payload = {
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
            "stream": False,
            "options": {"num_predict": max_tokens, "temperature": 0.7},
        }

        response = requests.post(url, json=payload, timeout=120)
        if response.status_code == 200:
            data = response.json()
            content = data.get("message", {}).get("content", "").strip()
            if content:
                return content
            return "[Ollama Error]: Respuesta vacía"
        else:
            return f"[Ollama Error {response.status_code}]: {response.text[:200]}"
    except requests.exceptions.ConnectionError:
        return "[Ollama Error]: No se pudo conectar a Ollama. Asegúrate de que esté ejecutándose."
    except Exception as e:
        return f"[Ollama Error]: {str(e)[:200]}"


def _call_with_fallback(prompt: str, api_key: str = "") -> str:
    """
    Intenta generar contenido con Gemini, y si falla, usa Ollama.
    Esta es la función que se llama desde agata_create.
    """
    # 1. Intentar Gemini
    result = _call_gemini(prompt, api_key)
    if not result.startswith("[Gemini Error]"):
        return result
    
    # 2. Fallback a Ollama
    print(f"[Agata] Gemini falló: {result[:100]}. Usando Ollama...")
    result = _call_ollama(prompt)
    if not result.startswith("[Ollama Error]"):
        return result
    
    # 3. Si todo falla, devolver el error
    return result


# =====================================================================
# PALETAS DE COLORES
# =====================================================================

PALETTES = {
    "elegant": {
        "primary": "#2C3E50", "secondary": "#E74C3C", "accent": "#3498DB",
        "light": "#ECF0F1", "dark": "#1A252F", "text": "#2C3E50",
    },
    "pastel": {
        "primary": "#6C5CE7", "secondary": "#FD79A8", "accent": "#00CEC9",
        "light": "#F8F5FF", "dark": "#2D1B69", "text": "#4A3F6B",
    },
    "corporativo": {
        "primary": "#1B4F72", "secondary": "#2874A6", "accent": "#D4AC0D",
        "light": "#EBF5FB", "dark": "#0B2D47", "text": "#1B4F72",
    },
    "moderno": {
        "primary": "#0A0A0A", "secondary": "#FF6B6B", "accent": "#4ECDC4",
        "light": "#F7F7F7", "dark": "#0A0A0A", "text": "#2D2D2D",
    },
    "rosa": {
        "primary": "#E91E63", "secondary": "#9C27B0", "accent": "#FFD54F",
        "light": "#FCE4EC", "dark": "#880E4F", "text": "#4A0072",
    },
    "nordico": {
        "primary": "#5E81AC", "secondary": "#BF616A", "accent": "#A3BE8C",
        "light": "#ECEFF4", "dark": "#2E3440", "text": "#3B4252",
    },
}


def _hex_to_rgb(h: str) -> tuple:
    """Convierte un color hexadecimal a tupla RGB."""
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


# =====================================================================
# FUNCIONES DE DISEÑO (Word y PowerPoint) - SIN CAMBIOS
# =====================================================================

def _build_table(doc, lines: list[str], pal: dict):
    """Construye una tabla en el documento Word."""
    try:
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import nsdecls
        from docx.oxml import parse_xml
    except ImportError:
        print("[Agata] python-docx no está completamente instalado.")
        return

    primary_rgb = _hex_to_rgb(pal["primary"])
    text_rgb = _hex_to_rgb(pal["text"])
    light_rgb = _hex_to_rgb(pal["light"])

    rows_data = []
    for line in lines:
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            rows_data.append(cells)

    if not rows_data:
        return

    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row_data in enumerate(rows_data):
        for col_idx in range(num_cols):
            cell = table.cell(row_idx, col_idx)
            cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(cell_text)

            if row_idx == 0:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(255, 255, 255)
                shading = parse_xml(
                    f'<w:shd {nsdecls("w")} w:fill="{pal["primary"].lstrip("#")}" w:val="clear"/>'
                )
                cell._tc.get_or_add_tcPr().append(shading)
            else:
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(*text_rgb)
                if row_idx % 2 == 0:
                    shading = parse_xml(
                        f'<w:shd {nsdecls("w")} w:fill="{pal["light"].lstrip("#")}" w:val="clear"/>'
                    )
                    cell._tc.get_or_add_tcPr().append(shading)

            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph("")


def _parse_slides_from_content(content: str) -> list[dict]:
    """Extrae diapositivas del contenido generado por la IA."""
    slides = []
    blocks = content.split("\n---\n")
    if len(blocks) == 1:
        blocks = content.split("\n\n##")

    for block in blocks:
        block = block.strip()
        if not block:
            continue

        lines = block.split("\n")
        slide_title = ""
        keyword = ""
        slide_lines = []

        for i, line in enumerate(lines):
            stripped = line.strip()
            if not stripped:
                continue

            if i == 0 and (stripped.startswith("## ") or stripped.startswith("# ")):
                raw = re.sub(r"^#+\s*", "", stripped)
                kw_match = re.search(r"\[keyword:\s*(.+?)\]", raw, re.IGNORECASE)
                if kw_match:
                    keyword = kw_match.group(1).strip()
                    raw = re.sub(r"\s*\[keyword:.*?\]", "", raw, flags=re.IGNORECASE).strip()
                slide_title = raw
                slide_lines.append(f"## {raw}")
            else:
                slide_lines.append(stripped)

        if not slide_title:
            for line in lines:
                stripped = line.strip()
                if stripped and not stripped.startswith("-") and not stripped.startswith("**"):
                    slide_title = stripped[:80]
                    break

        if not keyword and slide_title:
            keyword = f"{slide_title} high quality background"

        slide_text = "\n".join(slide_lines) if slide_lines else block
        slides.append({"title": slide_title, "keyword": keyword, "text": slide_text})

    return slides


# Placeholder para búsqueda de imágenes (se puede ampliar con Unsplash API)
def _buscar_y_descargar_imagen(keyword: str):
    """
    Busca y descarga una imagen desde internet (placeholder).
    En una implementación real, aquí se conectaría a Unsplash, Pexels, etc.
    """
    # Por ahora, siempre devuelve None para que no falle
    return None


def _analizar_paleta_y_brillo(img_path: str, pal: dict) -> dict:
    """Analiza la paleta de colores y brillo de una imagen (placeholder)."""
    # Valores predeterminados para que el diseño no falle
    return {
        "accent": _hex_to_rgb(pal["accent"]),
        "text_main": _hex_to_rgb(pal["light"]),
        "text_secondary": _hex_to_rgb(pal["text"]),
        "is_background_light": False,
    }


def _design_word(content: str, title: str, palette: str, player=None, speak=None) -> str:
    """Diseña un documento Word profesional."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import nsdecls
        from docx.oxml import parse_xml
    except ImportError:
        raise ImportError("python-docx no está instalado. Ejecuta: pip install python-docx")

    _ensure_folder()
    pal = PALETTES.get(palette, PALETTES["elegant"])
    primary_rgb = _hex_to_rgb(pal["primary"])
    secondary_rgb = _hex_to_rgb(pal["secondary"])
    accent_rgb = _hex_to_rgb(pal["accent"])
    text_rgb = _hex_to_rgb(pal["text"])

    filename = f"Agata_{title.replace(' ', '_')[:40]}_{uuid.uuid4().hex[:6]}.docx"
    filepath = AGATA_FOLDER / filename

    doc = Document()

    # Configurar estilos por defecto
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(*text_rgb)
    pf = style.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.3

    # Configurar estilos de títulos
    for i, (h_style, h_size, h_color) in enumerate([
        ("Heading 1", 22, primary_rgb),
        ("Heading 2", 16, secondary_rgb),
        ("Heading 3", 13, accent_rgb),
    ]):
        try:
            hs = doc.styles[h_style]
            hs.font.size = Pt(h_size)
            hs.font.bold = True
            hs.font.color.rgb = RGBColor(*h_color)
            hs.font.name = "Calibri Light"
            hs.paragraph_format.space_before = Pt(18 if i < 2 else 12)
            hs.paragraph_format.space_after = Pt(8)
        except Exception:
            pass

    # Márgenes
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.8)

    # ===== PORTADA =====
    for _ in range(5):
        doc.add_paragraph("")

    # Línea decorativa superior
    line_p = doc.add_paragraph()
    line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = line_p.add_run("━" * 40)
    lr.font.color.rgb = RGBColor(*accent_rgb)
    lr.font.size = Pt(10)

    # Título principal
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run(title.upper())
    tr.font.size = Pt(34)
    tr.font.bold = True
    tr.font.color.rgb = RGBColor(*primary_rgb)
    tr.font.name = "Calibri Light"

    # Subtítulo
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle_p.add_run("Diseñado por Agata  |  JARVIS Industries")
    sr.font.size = Pt(12)
    sr.font.color.rgb = RGBColor(*secondary_rgb)
    sr.font.italic = True

    # Fecha
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dr = date_p.add_run(time.strftime("%d de %B de %Y"))
    dr.font.size = Pt(10)
    dr.font.color.rgb = RGBColor(150, 150, 150)

    # Línea decorativa inferior
    line_p2 = doc.add_paragraph()
    line_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr2 = line_p2.add_run("━" * 40)
    lr2.font.color.rgb = RGBColor(*accent_rgb)
    lr2.font.size = Pt(10)

    doc.add_page_break()

    # ===== CONTENIDO =====
    lines = content.strip().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        # Detectar encabezados
        if line.startswith("# ") or line.startswith("## ") or line.startswith("### "):
            level = len(line) - len(line.lstrip("#"))
            heading_text = re.sub(r"^#+\s*", "", line)
            p = doc.add_paragraph()
            p.style = doc.styles[f"Heading {min(level, 3)}"]
            run = p.add_run(heading_text)

        # Insertar imagen (placeholder)
        elif line.startswith("[IMAGE:") or line.startswith("[IMAGEN:"):
            kw_match = re.search(r"\[(?:IMAGE|IMAGEN):\s*(.+?)\]", line, re.IGNORECASE)
            if kw_match:
                keyword = kw_match.group(1).strip()
                if player:
                    player.write_log(f"[Agata] Buscando imagen: {keyword}")
                img_path = _buscar_y_descargar_imagen(keyword)
                if img_path:
                    try:
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run()
                        run.add_picture(img_path, width=Inches(4.5))
                        doc.add_paragraph("")
                    except Exception:
                        pass

        # Insertar tabla
        elif line.startswith("[TABLE]") or line.startswith("[TABLA]"):
            table_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("[ENDTABLE]") and not lines[i].strip().startswith("[FIN_TABLA]"):
                table_lines.append(lines[i].strip())
                i += 1
            if table_lines:
                _build_table(doc, table_lines, pal)

        # Viñetas
        elif line.startswith("- "):
            bullet_text = line[2:]
            p = doc.add_paragraph(style="List Bullet")
            p.clear()
            run = p.add_run(bullet_text)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(*text_rgb)

        # Lista numerada
        elif re.match(r"^\d+[\.\)]\s", line):
            p = doc.add_paragraph(style="List Number")
            p.clear()
            run = p.add_run(re.sub(r"^\d+[\.\)]\s*", "", line))
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(*text_rgb)

        # Cita
        elif line.startswith("> "):
            quote_text = line[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.5)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(quote_text)
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = RGBColor(*secondary_rgb)

        # Separador
        elif line.startswith("---") or line.startswith("==="):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("─" * 50)
            r.font.color.rgb = RGBColor(*accent_rgb)
            r.font.size = Pt(8)

        # Párrafo normal
        else:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(*text_rgb)

        i += 1

    # ===== PIE DE PÁGINA =====
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("─" * 40)
    r.font.color.rgb = RGBColor(*accent_rgb)
    r.font.size = Pt(8)

    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_para.add_run("Creado con ❤ por Agata  |  JARVIS Industries")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(*secondary_rgb)
    fr.font.italic = True

    # Números de página
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fpr = fp.add_run()
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        fpr._r.append(fldChar1)
        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        fpr._r.append(instrText)
        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        fpr._r.append(fldChar2)
        fpr.font.size = Pt(8)
        fpr.font.color.rgb = RGBColor(150, 150, 150)

    doc.save(str(filepath))
    return str(filepath)


def _design_ppt(content: str, title: str, palette: str, player=None, speak=None) -> str:
    """Diseña una presentación PowerPoint profesional."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise ImportError("python-pptx no está instalado. Ejecuta: pip install python-pptx")

    def _overlay(slide, prs):
        from pptx.oxml.ns import qn
        ov = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        ov.fill.solid()
        ov.fill.fore_color.rgb = RGBColor(0, 0, 0)
        ov.line.fill.background()
        sp_el = ov._element
        sf = sp_el.find('.//' + qn('a:solidFill'))
        if sf is not None:
            srgb = sf.find(qn('a:srgbClr'))
            if srgb is not None:
                alpha_attr = qn('a:alpha')
                srgb.set(alpha_attr, '50000')
        else:
            srgb = sp_el.find('.//' + qn('a:srgbClr'))
            if srgb is not None:
                srgb.set(qn('a:alpha'), '50000')
        return ov

    _ensure_folder()
    pal = PALETTES.get(palette, PALETTES["moderno"])
    filename = f"Agata_{title.replace(' ', '_')[:40]}_{uuid.uuid4().hex[:6]}.pptx"
    filepath = AGATA_FOLDER / filename

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides_data = _parse_slides_from_content(content)
    if not slides_data:
        slides_data = [{"title": title, "keyword": f"{title} technology abstract", "text": content}]

    # ===== DIAPOSITIVA DE TÍTULO =====
    title_keyword = f"{title} abstract technology background"
    title_image = _buscar_y_descargar_imagen(title_keyword)
    title_info = _analizar_paleta_y_brillo(title_image, pal)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if title_image:
        slide.shapes.add_picture(title_image, 0, 0, prs.slide_width, prs.slide_height)

    if not title_image or title_info["is_background_light"]:
        _overlay(slide, prs)

    # Título
    txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2), Inches(10), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*title_info["text_main"])
    p.font.name = "Calibri Light"
    p.alignment = PP_ALIGN.CENTER

    # Subtítulo
    sub_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(10), Inches(1))
    sub_tf = sub_box.text_frame
    sp = sub_tf.paragraphs[0]
    sp.text = "Diseñado por Agata"
    sp.font.size = Pt(18)
    sp.font.color.rgb = RGBColor(*title_info["accent"])
    sp.font.italic = True
    sp.alignment = PP_ALIGN.CENTER

    # Línea decorativa
    line_shape = slide.shapes.add_shape(1, Inches(5), Inches(3.5), Inches(3), Inches(0.03))
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = RGBColor(*title_info["accent"])
    line_shape.line.fill.background()

    # ===== DIAPOSITIVAS DE CONTENIDO =====
    for slide_idx, sd in enumerate(slides_data):
        if player:
            player.write_log(f"[Agata] Generando diapositiva {slide_idx+1}: {sd['title'][:40]}")

        img_path = _buscar_y_descargar_imagen(sd["keyword"])
        info_color = _analizar_paleta_y_brillo(img_path, pal)

        slide = prs.slides.add_slide(prs.slide_layouts[6])

        if img_path:
            slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)

        has_text_color = info_color["text_main"]

        if not img_path or info_color["is_background_light"]:
            _overlay(slide, prs)
            if not img_path:
                has_text_color = _hex_to_rgb(pal["light"])
                info_color["accent"] = _hex_to_rgb(pal["accent"])
                info_color["text_secondary"] = (180, 180, 180)

        lines = sd["text"].strip().split("\n")
        y_pos = 1.0 if len(lines) < 6 else 0.5

        has_title = False
        for line in lines:
            line = line.strip()
            if not line:
                y_pos += 0.3
                continue

            if (line.startswith("## ") or line.startswith("# ")) and not has_title:
                heading_text = re.sub(r"^#+\s*", "", line)
                heading_text = re.sub(r"\s*\[keyword:.*?\]", "", heading_text, flags=re.IGNORECASE).strip()
                tb = slide.shapes.add_textbox(Inches(1), Inches(y_pos), Inches(11.3), Inches(1))
                ttf = tb.text_frame
                tp = ttf.paragraphs[0]
                tp.text = heading_text
                tp.font.size = Pt(32)
                tp.font.bold = True
                tp.font.color.rgb = RGBColor(*info_color["accent"])
                tp.font.name = "Calibri Light"
                tp.alignment = PP_ALIGN.LEFT
                y_pos += 1.2
                has_title = True

            elif line.startswith("- "):
                bullet = line[2:]
                tb = slide.shapes.add_textbox(Inches(1.5), Inches(y_pos), Inches(10.3), Inches(0.5))
                ttf = tb.text_frame
                tp = ttf.paragraphs[0]
                tp.text = f"   {bullet}"
                tp.font.size = Pt(16)
                tp.font.color.rgb = RGBColor(*has_text_color)
                tp.font.name = "Calibri"
                y_pos += 0.55

            elif line.startswith("**") and line.endswith("**"):
                highlight = line.strip("*")
                tb = slide.shapes.add_textbox(Inches(1), Inches(y_pos), Inches(11.3), Inches(0.6))
                ttf = tb.text_frame
                tp = ttf.paragraphs[0]
                tp.text = highlight
                tp.font.size = Pt(20)
                tp.font.bold = True
                tp.font.color.rgb = RGBColor(*info_color["accent"])
                tp.font.name = "Calibri"
                tp.alignment = PP_ALIGN.LEFT
                y_pos += 0.8

            else:
                tb = slide.shapes.add_textbox(Inches(1), Inches(y_pos), Inches(11.3), Inches(0.5))
                ttf = tb.text_frame
                tp = ttf.paragraphs[0]
                tp.text = line[:200]
                tp.font.size = Pt(16)
                tp.font.color.rgb = RGBColor(*has_text_color)
                tp.font.name = "Calibri"
                y_pos += 0.55

        # Número de diapositiva
        num_box = slide.shapes.add_textbox(Inches(12), Inches(7), Inches(1), Inches(0.4))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = str(slide_idx + 2)
        np.font.size = Pt(10)
        np.font.color.rgb = RGBColor(*info_color["text_secondary"])
        np.alignment = PP_ALIGN.RIGHT

    # ===== DIAPOSITIVA DE CIERRE =====
    closing_keyword = "thank you elegant background"
    closing_image = _buscar_y_descargar_imagen(closing_keyword)
    closing_info = _analizar_paleta_y_brillo(closing_image, pal)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if closing_image:
        slide.shapes.add_picture(closing_image, 0, 0, prs.slide_width, prs.slide_height)

    if not closing_image or closing_info["is_background_light"]:
        _overlay(slide, prs)

    thanks_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10), Inches(2))
    tf = thanks_box.text_frame
    tp = tf.paragraphs[0]
    tp.text = "Gracias"
    tp.font.size = Pt(52)
    tp.font.bold = True
    tp.font.color.rgb = RGBColor(*closing_info["text_main"])
    tp.font.name = "Calibri Light"
    tp.alignment = PP_ALIGN.CENTER

    agata_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.5), Inches(10), Inches(1))
    atf = agata_box.text_frame
    ap = atf.paragraphs[0]
    ap.text = "Diseñado con amor por Agata  |  JARVIS Industries"
    ap.font.size = Pt(16)
    ap.font.color.rgb = RGBColor(*closing_info["accent"])
    ap.font.italic = True
    ap.alignment = PP_ALIGN.CENTER

    prs.save(str(filepath))
    return str(filepath)


# =====================================================================
# FUNCIÓN PRINCIPAL EXPORTABLE (con fallback mejorado)
# =====================================================================

def agata_create(parameters: dict, player=None, speak=None):
    """
    Crea un documento Word o presentación PowerPoint con diseño profesional.

    Parámetros:
        type (str): "word" o "ppt"
        title (str): Título del documento
        topic (str): Tema o contenido a desarrollar
        palette (str): Nombre de la paleta de colores (elegant, pastel, corporativo, moderno, rosa, nordico)
        api_key (str): Clave API de Gemini (opcional, si no se proporciona usa la de config)

    Retorna:
        str: Mensaje de resultado con la ruta del archivo generado.
    """
    doc_type = parameters.get("type", "word").lower()
    title = parameters.get("title", "Documento Agata")
    topic = parameters.get("topic", "")
    palette = parameters.get("palette", "elegant")
    api_key = parameters.get("api_key", "")

    if doc_type not in ("word", "ppt"):
        doc_type = "word"

    if not topic:
        return "Necesito un tema (topic) para crear el documento. Por favor, dime de qué quieres que trate."

    if player:
        player.write_log(f"[Agata] Creando {doc_type.upper()}: {title}")

    # Construir prompt según el tipo de documento
    prompt = (
        f"Crea contenido profesional y bien estructurado para un {'documento Word' if doc_type == 'word' else 'presentación PowerPoint'} "
        f"titulado '{title}' sobre el tema: {topic}.\n\n"
    )

    if doc_type == "word":
        prompt += (
            "Formato del contenido (usa estos marcadores para elementos especiales):\n\n"
            "# Título Principal\n"
            "Introducción elegante y profesional (2-3 párrafos)\n\n"
            "## Sección Principal\n"
            "Contenido detallado con datos relevantes y análisis profundo\n"
            "- Punto clave 1 con detalle\n"
            "- Punto clave 2 con detalle\n\n"
            "Para insertar una imagen usa: [IMAGE: keyword en inglés]\n"
            "Ejemplo: [IMAGE: data analysis dashboard]\n\n"
            "Para insertar una tabla usa:\n"
            "[TABLE]\n"
            "| Columna 1 | Columna 2 | Columna 3 |\n"
            "| Dato 1 | Dato 2 | Dato 3 |\n"
            "| Dato 4 | Dato 5 | Dato 6 |\n"
            "[ENDTABLE]\n\n"
            "Para citas destacadas usa: > Texto de la cita\n"
            "Para separadores de sección usa: ---\n\n"
            "### Sub-sección\n"
            "Detalles complementarios\n\n"
            "## Conclusión\n"
            "Resumen elegante y cierre profesional\n\n"
            "Incluye al menos 1 tabla con datos relevantes y 1 imagen si el tema lo permite.\n"
            "Escribe todo en español, lenguaje profesional pero cálido.\n"
            "IMPORTANTE: No incluyas tu proceso de pensamiento ni análisis. "
            "Empieza directamente con el contenido del documento."
        )
    else:
        prompt += (
            "Formato: crea contenido por diapositivas separadas por '---'.\n"
            "Para cada diapositiva:\n"
            "## Título de Diapositiva [keyword: términos de búsqueda en inglés]\n"
            "**Idea Principal**\n"
            "- Punto clave 1\n"
            "- Punto clave 2\n"
            "- Punto clave 3\n"
            "Texto complementario breve\n\n"
            "IMPORTANTE: Después de cada título de diapositiva, incluye SIEMPRE [keyword: ...] "
            "con 2-4 palabras clave en INGLÉS que describan una imagen de fondo ideal para esa diapositiva. "
            "Ejemplo: ## Exploración Espacial [keyword: mars rover space exploration]\n"
            "Incluye 5-8 diapositivas. Escribe todo en español, lenguaje profesional pero cálido.\n"
            "IMPORTANTE: No incluyas tu proceso de pensamiento ni análisis. "
            "Empieza directamente con el contenido de la primera diapositiva."
        )

    # ===== GENERACIÓN DE CONTENIDO (USANDO FALLBACK MEJORADO) =====
    content = _call_with_fallback(prompt, api_key)

    # Si el contenido empieza con error, devolver el mensaje
    if content.startswith("[Gemini Error]") or content.startswith("[Ollama Error]"):
        return f"No pude generar el contenido. Error: {content}"

    # Si el contenido es demasiado corto o vacío
    if not content or len(content.strip()) < 20:
        return "El contenido generado está vacío o es demasiado corto."

    # ===== DISEÑO DEL DOCUMENTO =====
    if player:
        player.write_log(f"[Agata] Diseñando {doc_type.upper()} con paleta {palette}...")

    try:
        if doc_type == "word":
            filepath = _design_word(content, title, palette, player, speak)
        else:
            filepath = _design_ppt(content, title, palette, player, speak)
    except Exception as e:
        return f"Error al crear el archivo: {str(e)[:200]}"

    # ===== RESULTADO FINAL =====
    if os.path.exists(filepath):
        size_kb = os.path.getsize(filepath) / 1024
        result = (
            f"Tu {'documento Word' if doc_type == 'word' else 'presentación PowerPoint'} "
            f"'{title}' está listo!\n"
            f"Archivo: {Path(filepath).name}\n"
            f"Tamaño: {size_kb:.1f} KB\n"
            f"Paleta: {palette}\n"
            f"Ubicación: {AGATA_FOLDER}"
        )
        if player:
            player.write_log(f"[Agata] Listo: {Path(filepath).name} ({size_kb:.1f} KB)")

        # Abrir el archivo automáticamente
        try:
            if sys.platform == "win32":
                os.startfile(filepath)
            elif sys.platform == "darwin":
                subprocess.run(["open", filepath])
            else:
                subprocess.run(["xdg-open", filepath])
        except Exception:
            pass

        return result
    else:
        return f"El archivo se creó pero no pude encontrarlo en: {filepath}"


def list_palettes(parameters=None, player=None, speak=None):
    """Lista todas las paletas de colores disponibles."""
    result = "Paletas de diseño disponibles:\n\n"
    for name, pal in PALETTES.items():
        result += f"[{name.upper()}]\n"
        result += f"  Primario:  {pal['primary']}\n"
        result += f"  Secundario: {pal['secondary']}\n"
        result += f"  Acento:    {pal['accent']}\n\n"
    return result